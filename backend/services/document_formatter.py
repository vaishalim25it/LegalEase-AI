from io import BytesIO

import html
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from fpdf import FPDF


BRAND_NAME = "LegalEase"


def sanitize_text(text: str) -> str:
    """
    Convert text into a safe format for
    DOCX, PDF and HTML output.
    """

    if text is None:
        return ""

    text = str(text)

    replacements = {
        "\u2018": "'",
        "\u2019": "'",
        "\u201c": '"',
        "\u201d": '"',
        "\u2013": "-",
        "\u2014": "-",
        "\u00a0": " ",
    }

    for old, new in replacements.items():
        text = text.replace(old, new)

    return text


def is_heading(line: str) -> bool:
    """
    Detect common legal document headings.
    """

    line = line.strip()

    if not line:
        return False

    if re.match(r"^\d+[\.\)]\s+", line):
        return True

    if line.isupper() and len(line) < 100:
        return True

    return False


def format_docx(
    content: str,
    document_type: str = "Legal Document"
) -> bytes:
    """
    Create a DOCX file from generated legal document text.
    """

    document = Document()

    section = document.sections[0]

    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    # -------------------------
    # Brand
    # -------------------------

    brand = document.add_paragraph()

    brand.alignment = WD_ALIGN_PARAGRAPH.CENTER

    brand_run = brand.add_run(
        BRAND_NAME
    )

    brand_run.bold = True
    brand_run.font.name = "Times New Roman"
    brand_run.font.size = Pt(16)

    # -------------------------
    # Document Title
    # -------------------------

    title = document.add_paragraph()

    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    title_run = title.add_run(
        sanitize_text(
            document_type
        ).upper()
    )

    title_run.bold = True
    title_run.font.name = "Times New Roman"
    title_run.font.size = Pt(14)

    document.add_paragraph()

    # -------------------------
    # Main Content
    # -------------------------

    lines = sanitize_text(
        content
    ).splitlines()

    for line in lines:

        if not line.strip():

            document.add_paragraph()

            continue

        paragraph = document.add_paragraph()

        if is_heading(line):

            run = paragraph.add_run(
                line.strip()
            )

            run.bold = True
            run.font.name = "Times New Roman"
            run.font.size = Pt(11)

        else:

            run = paragraph.add_run(
                line.strip()
            )

            run.font.name = "Times New Roman"
            run.font.size = Pt(11)

    # -------------------------
    # Footer
    # -------------------------

    footer = section.footer

    footer_paragraph = (
        footer.paragraphs[0]
    )

    footer_paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    footer_run = footer_paragraph.add_run(
        "LegalEase - AI-generated drafting aid | "
        "Legal review recommended"
    )

    footer_run.font.name = (
        "Times New Roman"
    )

    footer_run.font.size = Pt(8)

    # -------------------------
    # Save DOCX
    # -------------------------

    output = BytesIO()

    document.save(output)

    return output.getvalue()


class LegalPDF(FPDF):

    def __init__(
        self,
        document_type="Legal Document"
    ):

        super().__init__()

        self.document_type = (
            document_type
        )

    def header(self):

        self.set_font(
            "Helvetica",
            "B",
            15
        )

        self.cell(
            0,
            10,
            BRAND_NAME,
            align="C"
        )

        self.ln(5)

        self.set_font(
            "Helvetica",
            "B",
            11
        )

        self.cell(
            0,
            8,
            sanitize_text(
                self.document_type
            ),
            align="C"
        )

        self.ln(8)

    def footer(self):

        self.set_y(-15)

        self.set_font(
            "Helvetica",
            "",
            8
        )

        self.cell(
            0,
            10,
            "LegalEase - AI-generated drafting aid | "
            "Legal review recommended",
            align="C"
        )


def format_pdf(
    content: str,
    document_type: str = "Legal Document"
) -> bytes:
    """
    Create a PDF file from generated legal
    document text.

    Handles long unbroken text safely.
    """

    pdf = LegalPDF(
        document_type
    )

    pdf.set_auto_page_break(
        auto=True,
        margin=20
    )

    pdf.add_page()

    pdf.set_font(
        "Helvetica",
        "",
        11
    )

    # Actual usable page width.
    page_width = pdf.epw

    lines = sanitize_text(
        content
    ).splitlines()

    for original_line in lines:

        line = original_line.strip()

        # -------------------------
        # Empty line
        # -------------------------

        if not line:

            pdf.ln(4)

            continue

        # -------------------------
        # Break long words
        # -------------------------

        words = line.split()

        safe_words = []

        for word in words:

            if len(word) > 30:

                chunks = [
                    word[i:i + 25]
                    for i in range(
                        0,
                        len(word),
                        25
                    )
                ]

                safe_words.extend(
                    chunks
                )

            else:

                safe_words.append(
                    word
                )

        line = " ".join(
            safe_words
        )

        # -------------------------
        # Reset cursor
        # -------------------------

        pdf.set_x(
            pdf.l_margin
        )

        # -------------------------
        # Heading
        # -------------------------

        if is_heading(line):

            pdf.set_font(
                "Helvetica",
                "B",
                11
            )

            pdf.multi_cell(
                page_width,
                7,
                line
            )

            pdf.set_font(
                "Helvetica",
                "",
                11
            )

        # -------------------------
        # Normal paragraph
        # -------------------------

        else:

            pdf.multi_cell(
                page_width,
                6,
                line
            )

    return bytes(
        pdf.output()
    )


def format_txt(
    content: str
) -> bytes:
    """
    Convert document content into TXT format.
    """

    return sanitize_text(
        content
    ).encode(
        "utf-8"
    )


def format_html_preview(
    content: str
) -> str:
    """
    Convert generated document text
    into an HTML preview.
    """

    safe_content = html.escape(
        sanitize_text(
            content
        )
    )

    safe_content = safe_content.replace(
        "\n",
        "<br>"
    )

    return f"""
    <div class="document-preview">

        <div class="document-brand">
            {BRAND_NAME}
        </div>

        <div class="document-content">
            {safe_content}
        </div>

    </div>
    """